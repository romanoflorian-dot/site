import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

d = json.load(open('data/obras.json')); im = json.load(open('data/imagens.json'))
F = "Arial"
TINTA, CINZA, PAPEL, VERM = "151515", "6B6762", "F1EFEA", "C8321E"
f_tit = Font(name=F, size=16, bold=True, color=TINTA)
f_sub = Font(name=F, size=10, color=CINZA)
f_cab = Font(name=F, size=10, bold=True, color="FFFFFF")
f_txt = Font(name=F, size=10, color=TINTA)
f_ok = Font(name=F, size=10, color="2E7D32", bold=True)
f_no = Font(name=F, size=10, color=VERM, bold=True)
fill_cab = PatternFill("solid", fgColor=TINTA)
fill_zebra = PatternFill("solid", fgColor=PAPEL)
fill_resp = PatternFill("solid", fgColor="FFF6CC")
linha = Border(bottom=Side(style="thin", color="CFCBC4"))
wrap = Alignment(wrap_text=True, vertical="top")
centro = Alignment(horizontal="center", vertical="top")

def cabecalho(ws, row, cols, larg):
    for i, (c, w) in enumerate(zip(cols, larg), 1):
        x = ws.cell(row=row, column=i, value=c); x.font = f_cab; x.fill = fill_cab
        x.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left")
        ws.column_dimensions[x.column_letter].width = w
    ws.row_dimensions[row].height = 28

wb = Workbook()
# ---------------------------------------------------------------- Obras
ws = wb.active; ws.title = "Obras"
ws["A1"] = "O que falta no site — por obra"; ws["A1"].font = f_tit
ws["A2"] = "florianoromano.art.br · ✓ já está no site · — falta · Preencha só a coluna amarela “Enviado?” à medida que subir o material."
ws["A2"].font = f_sub
cols = ["Obra", "Ano", "Tipo", "Fotos no site", "Áudio", "Vídeo", "Texto PT/EN", "Ficha técnica", "O que falta enviar", "Prioridade", "Pasta no Drive (site/…)", "Enviado?"]
cabecalho(ws, 5, cols, [26, 11, 17, 9, 8, 8, 10, 10, 52, 11, 26, 11])
r = 6
SEM_AUDIO_OK = {"desenho", "exposição", "vídeo"}
for o in sorted(d['obras'], key=lambda o: (-o['ano_num'], o['titulo'])):
    imgs = [i for i in (im.get(o['id'], []) or o.get('imagens', [])) if not i.get('aprovar')]
    aud = bool(o.get('audio') or o.get('soundcloud'))
    vid = bool(o.get('video')); txt = bool(o.get('texto')); fic = bool(o.get('materiais'))
    falta = []
    if not o['ano']: falta.append("ano da obra")
    if not imgs: falta.append("fotos (3 a 6)")
    if not aud and o['tipo'] not in SEM_AUDIO_OK: falta.append("áudio (MP3 ou link do SoundCloud)")
    if not txt: falta.append("texto seu ou de crítico sobre a obra")
    if not fic: falta.append("ficha técnica (materiais, dimensões, local)")
    if o['id'] == 'historias-anonimas': falta.append("confirmar duração")
    if o['id'] == 'esquinas': falta.append("confirmar suporte (tela ou madeira)")
    if o['id'] == 'cidade-sensivel': falta.append("links dos áudios do programa")
    if o['id'] == 'esquina': falta.append("fotos em alta (Pedro Tressi, até 5 MB)")
    vazia = not imgs and not txt
    if not falta: prio = "Completa"
    elif vazia or (o['ano_num'] >= 2016 and not imgs): prio = "Alta"
    elif not imgs or not txt: prio = "Média"
    else: prio = "Baixa"
    vals = [o['titulo'], o['ano'] or "?", o['tipo'], len(imgs) if imgs else "—", "✓" if aud else "—", "✓" if vid else "—",
            "✓" if txt else "—", "✓" if fic else "—", (lambda t: t[0].upper()+t[1:])("; ".join(falta)) if falta else "Nada", prio, o['id'] + "/", ""]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v); c.font = f_txt; c.border = linha
        c.alignment = centro if i in (4, 5, 6, 7, 8, 10, 12) else wrap
        if v == "✓": c.font = f_ok
        if v == "—": c.font = f_no
    ws.cell(row=r, column=1).font = Font(name=F, size=10, italic=True, color=TINTA)
    ws.cell(row=r, column=12).fill = fill_resp
    if r % 2 == 0:
        for i in range(1, 12): ws.cell(row=r, column=i).fill = fill_zebra
    r += 1
ult = r - 1
dv = DataValidation(type="list", formula1='"Sim,Parcial,Não tenho"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"L6:L{ult}")
for txt_, cor in [("Alta", "F6D5CF"), ("Média", "F3E6C4"), ("Baixa", "E4E1DB"), ("Completa", "D8EBD9")]:
    ws.conditional_formatting.add(f"J6:J{ult}", CellIsRule(operator="equal", formula=[f'"{txt_}"'], fill=PatternFill("solid", fgColor=cor)))
ws.conditional_formatting.add(f"L6:L{ult}", CellIsRule(operator="equal", formula=['"Sim"'], fill=PatternFill("solid", fgColor="D8EBD9")))
ws["A3"] = "Prioridade alta:"; ws["A3"].font = Font(name=F, size=10, bold=True)
ws["B3"] = f'=COUNTIF(J6:J{ult},"Alta")'; ws["B3"].font = f_txt
ws["C3"] = "Já enviadas:"; ws["C3"].font = Font(name=F, size=10, bold=True)
ws["D3"] = f'=COUNTIF(L6:L{ult},"Sim")&" de "&COUNTIF(J6:J{ult},"<>Completa")'; ws["D3"].font = f_txt
ws.freeze_panes = "B6"; ws.auto_filter.ref = f"A5:L{ult}"
ws.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Perguntas
wp = wb.create_sheet("Perguntas")
wp["A1"] = "Perguntas e dados a confirmar"; wp["A1"].font = f_tit
wp["A2"] = "Responda na coluna amarela. Uma palavra basta."; wp["A2"].font = f_sub
cabecalho(wp, 4, ["#", "Assunto", "Pergunta", "O que está no site hoje", "Sua resposta"], [5, 22, 58, 30, 36])
P = [
 ("Contato", "Qual e-mail deve aparecer na página Sobre?", "“e-mail a definir”"),
 ("Prêmio CCBB", "Errância recebeu o Prêmio CCBB Contemporâneo em 2015, 2016 ou 2017? As fontes divergem.", "2016"),
 ("HOBRA", "A residência HOBRA (Brasil–Holanda) foi em 2016 ou 2017?", "2017"),
 ("Esquinas I–X", "Os desenhos são acrílica sobre tela ou sobre madeira?", "conforme catálogo"),
 ("Histórias Anônimas", "A ficha diz 3:45; o áudio que está no site tem 2:12. É um trecho? Há a versão completa?", "áudio de 2:12"),
 ("Virada Sônica", "Em que ano foi a Virada Sônica?", "2024 (suposição)"),
 ("Espreguiçadeira siamesa", "“Espreguiçadeira sonora siamesa (2018)”, citada no PIPA, é obra separada? Entra no site?", "não está"),
 ("Vídeo ArtRio", "O vídeo das intervenções Bradesco/ArtRio mostra a obra de Maya Dikstein. Entra no site?", "não está"),
 ("Coletivas", "Quais obras suas estavam em Arte pra Sentir, Corpo de Som, Paquetá Experimenta e Permanências e Destruições?", "só o vídeo e o título"),
 ("Rádio Invisível", "Em que ano foi feito?", "sem ano"),
 ("Cidade Sensível", "Links dos áudios do programa (SoundCloud ou outro).", "sem áudio"),
 ("Fotos Pedro Tressi", "Versões em alta resolução (até 5 MB cada) das fotos de A Pele do Mundo.", "fotos extraídas do catálogo"),
 ("Vídeos antigos", "Se houver no YouTube, versões de Guarda-Chuva Sonoro, Sapatos Sonoros e RadioMochilas (os do Vimeo antigo saíram do ar).", "versões do Canal-Arte"),
 ("Inglês", "Quer que alguém revise os textos em inglês? (opcional)", "tradução feita por mim"),
 ("Tarefa sua", "No GitHub: Settings → Pages → marcar “Enforce HTTPS”, se ainda não estiver marcado.", "—"),
]
for i, (a, q, h) in enumerate(P, 1):
    row = 4 + i
    for j, v in enumerate([i, a, q, h, ""], 1):
        c = wp.cell(row=row, column=j, value=v); c.font = f_txt; c.alignment = wrap; c.border = linha
    wp.cell(row=row, column=2).font = Font(name=F, size=10, bold=True)
    wp.cell(row=row, column=4).font = f_sub
    wp.cell(row=row, column=5).fill = fill_resp
    wp.row_dimensions[row].height = 30
wp.freeze_panes = "A5"; wp.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Como enviar
wc = wb.create_sheet("Como enviar")
wc["A1"] = "Como enviar o material"; wc["A1"].font = f_tit
wc.column_dimensions["A"].width = 20; wc.column_dimensions["B"].width = 90
L = [
 ("Onde", "Na pasta “site” do Google Drive, crie uma subpasta por obra com o nome da coluna “Pasta no Drive” (ex.: site/mar-sobre-mar/)."),
 ("Fotos", "JPG, até 5 MB cada (arquivos maiores travam a transferência). 3 a 6 por obra. Lado maior entre 2000 e 3000 px é suficiente."),
 ("Crédito das fotos", "Coloque o nome do fotógrafo no nome do arquivo (ex.: mar-sobre-mar_foto-Wilton-Montenegro.jpg) ou num .txt na pasta."),
 ("Áudio", "MP3 (até 10 MB) ou o link do SoundCloud. Um trecho de 1 a 3 minutos basta."),
 ("Vídeo", "Só o link do YouTube ou Vimeo, público ou não listado."),
 ("Texto", "Pode ser um parágrafo bruto, um trecho de catálogo ou o nome do texto crítico. Eu reduzo e traduzo."),
 ("Ficha técnica", "Materiais, dimensões, local e ano da primeira montagem. Pode ir num .txt na pasta da obra."),
 ("Exemplo", "site/praca-sonora/ → praca-sonora_01_foto-Fulano.jpg, praca-sonora_02_foto-Fulano.jpg, ficha.txt, link-soundcloud.txt"),
 ("Obras de terceiros", "Fotos e áudios de outras pessoas só entram com crédito e depois da sua confirmação."),
]
for i, (a, b) in enumerate(L, 3):
    wc.cell(row=i, column=1, value=a).font = Font(name=F, size=10, bold=True)
    c = wc.cell(row=i, column=2, value=b); c.font = f_txt; c.alignment = wrap
    wc.row_dimensions[i].height = 30
wc.sheet_view.showGridLines = False
for w in wb.worksheets:
    w.page_setup.orientation = "landscape"; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
wb.save("entregas/o-que-falta-no-site.xlsx")
print("linhas", ult - 5)
